import docx

from .base import Document


def load_docx(path: str) -> list[Document]:
    documento = docx.Document(path)
    text = "\n".join(p.text for p in documento.paragraphs if p.text)
    return [Document(text, {"source": path, "type": "docx"})]
